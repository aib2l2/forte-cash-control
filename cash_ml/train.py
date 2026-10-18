"""Train from point labels, keeping validation and test videos out of training."""

import argparse
import importlib.metadata
import json
import time
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import balanced_accuracy_score, confusion_matrix, precision_score, recall_score
from sklearn.preprocessing import StandardScaler

from cash_ml.features import MODEL_DIR, PREPROCESS_VERSION, WEIGHTS, Extractor, sha256


def read_labels(path, videos):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = wb["labels"]
        rows = list(sheet.iter_rows(values_only=True))
        headers = [str(x).strip() if x is not None else "" for x in rows[0]]
        if not {"video", "violation", "timecode_sec", "type"}.issubset(headers):
            raise ValueError("Не найдены обязательные столбцы разметки")
        result = []
        for row_number, row in enumerate(rows[1:], start=2):
            record = dict(zip(headers, row, strict=True))
            name = record.get("video")
            if name is None:
                continue
            if record["violation"] not in (0, 1):
                raise ValueError(f"Некорректная метка: {name}")
            name = str(name).strip()
            file = videos / f"{name}.mov"
            if not file.is_file():
                raise ValueError(f"Нет видео: {file}")
            y = int(record["violation"])
            t = record.get("timecode_sec")
            if y and (not isinstance(t, (int, float)) or not np.isfinite(t) or t < 0):
                raise ValueError(f"Нет корректного таймкода: {name}")
            event_type = record.get("type")
            if y and (not isinstance(event_type, str) or not event_type.strip()):
                raise ValueError(f"Нет типа нарушения: {name}")
            if not y and event_type not in (None, ""):
                raise ValueError(f"Тип нарушения у нормального видео: {name}")
            result.append({"id": name, "row_number": row_number, "label": y, "timecode": t,
                           "type": event_type, "path": file, "sha256": sha256(file)})
        grouped = {}
        for record in result:
            digest = record["sha256"]
            if digest not in grouped:
                grouped[digest] = {**record, "aliases": [], "annotations": [], "source_rows": []}
            group = grouped[digest]
            if group["label"] != record["label"]:
                raise ValueError("Одинаковые видео имеют противоречивые бинарные метки")
            if record["id"] not in group["aliases"]:
                group["aliases"].append(record["id"])
            group["source_rows"].append({k: record[k] for k in ("row_number", "id", "label", "timecode", "type")})
            if record["label"]:
                annotation = {"timecode": record["timecode"], "type": record["type"]}
                if annotation not in group["annotations"]:
                    group["annotations"].append(annotation)
        return list(grouped.values())
    finally:
        wb.close()


def metrics(labels, scores, threshold):
    predicted = np.asarray(scores) >= threshold
    return {"videos": len(labels), "threshold": float(threshold),
            "balanced_accuracy": float(balanced_accuracy_score(labels, predicted)),
            "precision": float(precision_score(labels, predicted, zero_division=0)),
            "recall": float(recall_score(labels, predicted, zero_division=0)),
            "confusion_matrix_tn_fp_fn_tp": confusion_matrix(labels, predicted, labels=[0, 1]).tolist()}


def split_with_type_coverage(records, seed=42):
    """Split whole SHA groups; singleton event types must remain in training."""
    rng = np.random.default_rng(seed)
    positives = [r for r in records if r["label"]]
    negatives = [r for r in records if not r["label"]]
    rng.shuffle(positives)
    rng.shuffle(negatives)
    counts = Counter(t for r in positives for t in {a["type"] for a in r["annotations"]})
    wanted_positive = max(1, round(len(positives) * 0.2))
    wanted_negative = max(1, round(len(negatives) * 0.2))
    held = []
    for r in positives:
        types = {a["type"] for a in r["annotations"]}
        if len(held) < 2 * wanted_positive and all(counts[t] > 1 for t in types):
            held.append(r)
            counts.subtract(types)
    if len(held) != 2 * wanted_positive or len(negatives) <= 2 * wanted_negative:
        raise ValueError("Недостаточно групп для разделения с сохранением всех типов в train")
    held_hashes = {r["sha256"] for r in held}
    train = [r for r in positives if r["sha256"] not in held_hashes] + negatives[2*wanted_negative:]
    val = held[::2] + negatives[:wanted_negative]
    test = held[1::2] + negatives[wanted_negative:2*wanted_negative]
    return train, val, test


def annotation_examples(records, data):
    x, y, locations = [], [], []
    for r in records:
        item = data[r["id"]]
        for annotation in r["annotations"]:
            index = int(np.argmin(abs(item["centers"] - annotation["timecode"])))
            x.append(item["features"][index])
            y.append(annotation["type"])
            locations.append({"video": r["id"], "timecode": annotation["timecode"], "window_index": index})
    return np.asarray(x), np.asarray(y), locations


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--labels", type=Path, required=True)
    parser.add_argument("--videos", type=Path, required=True)
    args = parser.parse_args()
    records = read_labels(args.labels, args.videos)
    MODEL_DIR.mkdir(parents=True, exist_ok=True)
    cache = MODEL_DIR / "features"
    cache.mkdir(exist_ok=True)
    extractor = Extractor(download=True)
    data = {}
    for i, record in enumerate(records):
        began = time.monotonic()
        cached = cache / f"{record['sha256']}-{PREPROCESS_VERSION}.npz"
        if not cached.exists():
            np.savez_compressed(cached, **extractor.extract(record["path"]))
        with np.load(cached, allow_pickle=False) as saved:
            data[record["id"]] = {key: saved[key] for key in saved.files}
        duration = float(data[record["id"]]["duration"])
        if any(a["timecode"] >= duration for a in record["annotations"]):
            raise ValueError(f"Таймкод вне видео: {record['id']}")
        print(f"[{i + 1}/{len(records)}] {record['id']}: {duration:.1f}s, "
              f"{len(data[record['id']]['centers'])} windows, {time.monotonic()-began:.1f}s", flush=True)
    train, val, test = split_with_type_coverage(records)
    allowed_types = sorted({a["type"] for r in records for a in r["annotations"]})
    x, y, sample_weights = [], [], []
    class_counts = Counter(r["label"] for r in train)
    for record in train:
        item = data[record["id"]]
        if record["label"]:
            indices = sorted({int(np.argmin(abs(item["centers"] - a["timecode"])))
                              for a in record["annotations"]})
        else:
            indices = list(range(len(item["features"])))
        for index in indices:
            x.append(item["features"][index])
            y.append(record["label"])
            sample_weights.append(1 / (len(indices) * class_counts[record["label"]]))
    weights = np.asarray(sample_weights)
    weights *= len(weights) / weights.sum()
    scaler = StandardScaler().fit(x, sample_weight=weights)
    classifier = LogisticRegression(C=0.1, max_iter=3000, random_state=42)
    classifier.fit(scaler.transform(x), y, sample_weight=weights)

    # A second trained head predicts only exact strings present in the source XLSX.
    type_x, type_y, _ = annotation_examples(train, data)
    type_scaler = StandardScaler().fit(type_x)
    type_classifier = LogisticRegression(C=0.1, class_weight="balanced", max_iter=3000, random_state=42)
    type_classifier.fit(type_scaler.transform(type_x), type_y)
    if list(type_classifier.classes_) != allowed_types:
        raise ValueError("Не все типы из Excel представлены в обучении")

    def predict(items):
        return [float(classifier.predict_proba(scaler.transform(data[r["id"]]["features"]))[:, 1].max()) for r in items]

    val_scores = predict(val)
    val_labels = [r["label"] for r in val]
    candidates = np.r_[0.0, np.unique(val_scores), np.nextafter(max(val_scores), np.inf)]
    threshold = max(candidates, key=lambda t: (
        balanced_accuracy_score(val_labels, np.asarray(val_scores) >= t), -abs(t - 0.5)))
    test_scores = predict(test)

    def type_metrics(items):
        event_x, expected, locations = annotation_examples(items, data)
        predicted = type_classifier.predict(type_scaler.transform(event_x))
        detection_scores = classifier.predict_proba(scaler.transform(event_x))[:, 1]
        return {
            "annotation_windows": len(expected),
            "accuracy_at_annotated_windows": float(np.mean(expected == predicted)),
            "detected_and_correct_type_at_annotated_windows": int(np.sum((detection_scores >= threshold) & (expected == predicted))),
            "class_order": allowed_types,
            "confusion_matrix": confusion_matrix(expected, predicted, labels=allowed_types).tolist(),
            "per_type": {t: {"support": int(np.sum(expected == t)),
                             "correct": int(np.sum((expected == t) & (predicted == t)))} for t in allowed_types},
            "predictions": [{**location, "expected": str(actual), "predicted": str(pred), "detected": bool(score >= threshold)}
                            for location, actual, pred, score in zip(locations, expected, predicted, detection_scores, strict=True)],
        }

    report = {
        "format_version": 2,
        "created_utc": datetime.now(UTC).isoformat(),
        "versions": {name: importlib.metadata.version(name) for name in
                     ["torch", "torchvision", "scikit-learn", "numpy", "opencv-python-headless"]},
        "model": "frozen R3D-18 + binary LogisticRegression + 11-type LogisticRegression",
        "allowed_types": allowed_types,
        "type_train_counts": dict(Counter(type_y)),
        "preprocess_version": PREPROCESS_VERSION, "labels_sha256": sha256(args.labels),
        "weights_sha256": sha256(WEIGHTS), "threshold": float(threshold),
        "unique_video_label_counts": dict(Counter(r["label"] for r in records)),
        "source_row_count": sum(len(r["source_rows"]) for r in records),
        "duplicate_groups": [r["aliases"] for r in records if len(r["aliases"]) > 1],
        "split_method": "SHA256 groups, seed 42; every type retained in train; singleton types cannot be evaluated independently; scene groups NOT reviewed",
        "limitations": ["Point labels, not reviewed event intervals", "Scene leakage not ruled out",
                        "Singleton types have no independent test examples", "Scores are not calibrated probabilities",
                        "New split supersedes v1 metrics; evaluation is exploratory"],
        "validation": metrics(val_labels, val_scores, threshold),
        "test": metrics([r["label"] for r in test], test_scores, threshold),
        "type_validation": type_metrics(val),
        "type_test": type_metrics(test),
        "splits": {name: [{k: v for k, v in r.items() if k != "path"} for r in items]
                   for name, items in [("train", train), ("validation", val), ("test", test)]},
        "test_predictions": [{"video": r["id"], "label": r["label"], "score": score}
                             for r, score in zip(test, test_scores, strict=True)],
    }
    temporary = MODEL_DIR / "typed_classifier.tmp.npz"
    np.savez(temporary, mean=scaler.mean_, scale=scaler.scale_, coef=classifier.coef_[0],
             intercept=classifier.intercept_[0], threshold=threshold,
             type_mean=type_scaler.mean_, type_scale=type_scaler.scale_,
             type_coef=type_classifier.coef_, type_intercept=type_classifier.intercept_,
             type_classes=type_classifier.classes_.astype(str), format_version=2,
             preprocess_version=PREPROCESS_VERSION)
    temporary.replace(MODEL_DIR / "typed_classifier.npz")
    report["classifier_sha256"] = sha256(MODEL_DIR / "typed_classifier.npz")
    report_path = MODEL_DIR / "typed_report.tmp.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2))
    report_path.replace(MODEL_DIR / "typed_report.json")
    print(json.dumps({k: report[k] for k in ["allowed_types", "type_train_counts", "validation", "test", "type_test"]}, indent=2), flush=True)


if __name__ == "__main__":
    main()
